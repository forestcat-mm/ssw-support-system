"""登録支援機関の帳簿の出力（T57。ADR-0036・ADR-0044）。

- 外国人ごと: ③ 支援対象者に関する管理簿・④ 支援の実施に関する管理簿（画面・自作の xlsx を
  PDF に）。
- 機関全体: ① 支援実施体制に関する管理簿・② 支援の委託契約に関する管理簿・月ごとの集計
  （自作の xlsx）。
自作の Excel（入管の様式ではない）なので openpyxl で作る（制約3 の対象外）。出力は監査ログに
EXPORT。
"""

from __future__ import annotations

import datetime
import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]

from apps.contracts.models import EmploymentContract
from apps.core.models import EmploymentType, Office, Organization, Staff
from apps.filings.models import Filing, FilingKind, FilingStatus
from apps.forms_engine.models import GeneratedDocument, GenerationStatus
from apps.hosts.models import SupportContract, SupportScope
from apps.support.models import (
    Briefing,
    Consultation,
    SupportPlan,
    SupportRecord,
)
from apps.workers.models import ForeignWorker, LatestResidenceRecord
from services import books, form_signatures

# --- 支援の期間（ADR-0036 決定 3: 雇用契約と全部委託の支援委託契約が重なる期間）-----------------


@dataclass(frozen=True)
class SupportPeriod:
    host_name: str
    starts_on: datetime.date
    ends_on: datetime.date | None
    end_reason: str  # 終了の理由（雇用契約の終了・支援委託契約の終了。続いていれば空）


def _span(start: datetime.date, end: datetime.date | None) -> tuple[datetime.date, datetime.date]:
    return start, end or datetime.date.max


def support_periods(worker: ForeignWorker) -> list[SupportPeriod]:
    periods: list[SupportPeriod] = []
    contracts = EmploymentContract.objects.filter(worker=worker).select_related("host_organization")
    for contract in contracts.order_by("current_starts_on", "pk"):
        contract_start = (
            contract.versions.order_by("version").values_list("starts_on", flat=True).first()
            or contract.current_starts_on
        )
        contract_end = books.contract_end(contract)
        supports = SupportContract.objects.filter(
            host_organization=contract.host_organization, scope=SupportScope.FULL
        ).order_by("starts_on", "pk")
        for support in supports:
            support_end = support.terminated_on or support.ends_on
            low = max(_span(contract_start, contract_end)[0], support.starts_on)
            high = min(
                _span(contract_start, contract_end)[1], _span(support.starts_on, support_end)[1]
            )
            if low > high:
                continue
            end = None if high == datetime.date.max else high
            reason = ""
            if end is not None and contract_end is not None and end == contract_end:
                reason = (
                    contract.get_termination_reason_display()
                    if contract.terminated_on
                    else "雇用契約の期間満了"
                )
            elif end is not None:
                reason = (
                    support.get_termination_category_display()
                    if support.terminated_on
                    else "支援委託契約の期間満了"
                )
            periods.append(SupportPeriod(contract.host_organization.name, low, end, str(reason)))
    return periods


# --- 外国人ごとの帳簿（③④）------------------------------------------------------------


@dataclass(frozen=True)
class BookRecord:
    """④ の 1 行（支援実績・相談・定期面談）。"""

    occurred_on: datetime.date
    item: str
    method: str
    place: str
    staff: str
    position: str
    interpreter: str
    content: str
    declined: bool = False


@dataclass
class WorkerBook:
    worker: ForeignWorker
    card_number: str
    hosts: list[str]
    plan_versions: list[str]
    periods: list[SupportPeriod]
    retention_until: datetime.date | None
    records: list[BookRecord] = field(default_factory=list)
    documents: list[tuple[GeneratedDocument, str]] = field(default_factory=list)


def _summary(text: str, limit: int = 200) -> str:
    text = " ".join(text.split())
    return text if len(text) <= limit else text[: limit - 1] + "…"


def worker_book(worker: ForeignWorker) -> WorkerBook:
    residence = LatestResidenceRecord.objects.select_related("record").filter(worker=worker).first()
    contracts = EmploymentContract.objects.filter(worker=worker).select_related("host_organization")
    plans = SupportPlan.objects.filter(worker=worker).order_by("pk")
    plan_versions = [
        f"{plan.host_organization.name if plan.host_organization else ''} 第{version.version}版"
        f"（{version.applies_from:%Y-%m-%d}〜）"
        for plan in plans
        for version in plan.versions.order_by("version")
    ]
    book = WorkerBook(
        worker=worker,
        card_number=residence.record.card_number if residence else "",
        hosts=sorted({contract.host_organization.name for contract in contracts}),
        plan_versions=plan_versions,
        periods=support_periods(worker),
        retention_until=books.retention_until(worker),
    )
    records = (
        SupportRecord.objects.filter(worker=worker, voided_at__isnull=True)
        .select_related("staff")
        .order_by("occurred_on", "pk")
    )
    briefings = {b.record_id: b for b in Briefing.objects.filter(worker=worker)}
    consultations = {c.record_id: c for c in Consultation.objects.filter(worker=worker)}
    for record in records:
        details = record.details or {}
        briefing = briefings.get(record.pk)
        consultation = consultations.get(record.pk)
        place = str(details.get("place", "")) or (briefing.place if briefing else "")
        interpreter = (
            (briefing.interpreter if briefing else "")
            or (consultation.interpreter if consultation else "")
            or str(details.get("interpreter", ""))
        )
        content = record.body
        if consultation is not None:
            labels = f"{consultation.get_category_display()}・{consultation.get_status_display()}"
            content = f"{labels} {content}"
        book.records.append(
            BookRecord(
                occurred_on=record.occurred_on,
                item=str(record.get_kind_display()),
                method=str(record.get_method_display()),
                place=place,
                staff=record.staff.name,
                position=record.staff_position or record.staff.position,
                interpreter=interpreter,
                content=_summary(content),
                declined=record.declined_by_worker,
            )
        )
    documents = GeneratedDocument.objects.filter(
        worker=worker, status=GenerationStatus.COMPLETED
    ).select_related("template")
    book.documents = [
        (document, form_signatures.state(document))
        for document in documents.order_by("generated_at")
    ]
    return book


def _sheet(
    workbook: Workbook, title: str, header: Sequence[str], rows: Iterable[Sequence[Any]]
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(list(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(
            [value.isoformat() if isinstance(value, datetime.date) else value for value in row]
        )
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(10, width * 1.6), 60)
        for cell in column:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def worker_book_xlsx(book: WorkerBook) -> bytes:
    worker = book.worker
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(
        workbook,
        "③支援対象者",
        ["項目", "内容"],
        [
            ["氏名", worker.name_romaji],
            ["生年月日", worker.birth_date],
            ["国籍・地域", worker.nationality.name],
            ["性別", worker.get_sex_display()],
            ["在留カード番号", book.card_number],
            ["雇用する特定技能所属機関", "、".join(book.hosts)],
            ["支援計画", "\n".join(book.plan_versions)],
            *[
                [
                    "支援の期間",
                    f"{period.host_name} {period.starts_on:%Y-%m-%d}〜"
                    + (
                        f"{period.ends_on:%Y-%m-%d}（{period.end_reason}）"
                        if period.ends_on
                        else ""
                    ),
                ]
                for period in book.periods
            ],
            ["帳簿の保存義務の満了日（目安）", book.retention_until or "まだ決まらない"],
        ],
    )
    _sheet(
        workbook,
        "④支援の実施",
        [
            "実施日",
            "支援の項目",
            "方法",
            "場所",
            "担当者",
            "役職",
            "通訳人",
            "内容",
            "本人の申出により実施せず",
        ],
        [
            [
                record.occurred_on,
                record.item,
                record.method,
                record.place,
                record.staff,
                record.position,
                record.interpreter,
                record.content,
                "○" if record.declined else "",
            ]
            for record in book.records
        ],
    )
    _sheet(
        workbook,
        "④関係する書類",
        ["生成日", "様式", "版", "署名"],
        [
            [
                document.generated_at.date() if document.generated_at else "",
                f"{document.template.form_id} {document.template.name}",
                document.template.valid_from,
                state,
            ]
            for document, state in book.documents
        ],
    )
    return _bytes(workbook)


# --- 機関全体（①②と月ごとの集計）-------------------------------------------------------


def _months(start: datetime.date, end: datetime.date) -> list[datetime.date]:
    months = []
    day = start.replace(day=1)
    while day <= end:
        months.append(day)
        day = (day.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
    return months


def _month_end(day: datetime.date) -> datetime.date:
    return (day.replace(day=28) + datetime.timedelta(days=4)).replace(day=1) - datetime.timedelta(
        days=1
    )


@dataclass(frozen=True)
class MonthCount:
    month: datetime.date
    supported: int
    missing: int
    full_time: int
    part_time: int


def monthly_counts(start: datetime.date, end: datetime.date) -> list[MonthCount]:
    """月ごとの支援人数・行方不明者数・月末の職員数。

    支援人数はその月に 1 日でも支援の期間がある外国人、行方不明者数は届出タスク「行方不明」の
    起算日、職員数は今の有効な担当者を登録の日から数える（常勤・非常勤。推定）。
    """
    periods = {worker.pk: support_periods(worker) for worker in ForeignWorker.objects.all()}
    missing = list(
        Filing.objects.filter(kind=FilingKind.MISSING)
        .exclude(status=FilingStatus.CANCELLED)
        .values_list("basis_on", flat=True)
    )
    staff = list(Staff.objects.filter(is_active=True).values_list("created_at", "employment_type"))
    result = []
    for month in _months(start, end):
        last = _month_end(month)
        supported = sum(
            1
            for worker_periods in periods.values()
            if any(
                period.starts_on <= last and (period.ends_on is None or period.ends_on >= month)
                for period in worker_periods
            )
        )
        lost = sum(1 for day in missing if month <= day <= last)
        on_month = [kind for created, kind in staff if created.date() <= last]
        result.append(
            MonthCount(
                month,
                supported,
                lost,
                sum(1 for kind in on_month if kind == EmploymentType.FULL_TIME),
                sum(1 for kind in on_month if kind != EmploymentType.FULL_TIME),
            )
        )
    return result


def organization_book_xlsx(start: datetime.date, end: datetime.date) -> bytes:
    organization = Organization.objects.first()
    workbook = Workbook()
    workbook.remove(workbook.active)
    rows: list[list[Any]] = []
    if organization is not None:
        rows += [
            ["名称", organization.name],
            ["住所", f"〒{organization.postal_code} {organization.address}"],
            ["代表者", f"{organization.representative_title} {organization.representative_name}"],
            ["法人番号", organization.corporate_number],
            ["登録番号", organization.registration_number],
        ]
        for officer in organization.officers.filter(voided_at__isnull=True):
            rows.append(
                [
                    "役員",
                    f"{officer.position} {officer.name} {officer.address}"
                    + (f"（{officer.resigned_on:%Y-%m-%d} 退任）" if officer.resigned_on else ""),
                ]
            )
    for office in Office.objects.order_by("pk"):
        rows.append(
            [
                "事業所",
                f"{office.name} 〒{office.postal_code} {office.address} {office.phone_number}",
            ]
        )
    _sheet(workbook, "①支援実施体制", ["項目", "内容"], rows)
    _sheet(
        workbook,
        "①担当者",
        [
            "氏名",
            "役職",
            "事務所",
            "支援責任者",
            "支援担当者",
            "雇用形態",
            "言語",
            "講習の修了日",
            "有効",
        ],
        [
            [
                member.name,
                member.position,
                member.office.name,
                "○" if member.is_support_manager else "",
                "○" if member.is_support_staff else "",
                member.get_employment_type_display(),
                "、".join(language.name for language in member.languages.all()),
                member.training_completed_on or "",
                "○" if member.is_active else "",
            ]
            for member in Staff.objects.select_related("office").order_by("pk")
        ],
    )
    _sheet(
        workbook,
        "②支援の委託契約",
        [
            "受入れ機関",
            "締結日",
            "開始日",
            "終了日",
            "委託の範囲",
            "委託料（月額）",
            "内訳",
            "終了の区分",
        ],
        [
            [
                support.host_organization.name,
                support.concluded_on,
                support.starts_on,
                support.terminated_on or support.ends_on or "",
                support.get_scope_display(),
                support.monthly_fee or "",
                "、".join(f"{fee.name} {fee.amount:,}円" for fee in support.fees.all()),
                support.get_termination_category_display() if support.terminated_on else "",
            ]
            for support in SupportContract.objects.select_related("host_organization").order_by(
                "host_organization__name", "starts_on"
            )
        ],
    )
    _sheet(
        workbook,
        "月ごとの集計",
        ["月", "支援人数", "行方不明者数", "職員数（常勤）", "職員数（非常勤）"],
        [
            [
                f"{count.month:%Y-%m}",
                count.supported,
                count.missing,
                count.full_time,
                count.part_time,
            ]
            for count in monthly_counts(start, end)
        ],
    )
    return _bytes(workbook)
