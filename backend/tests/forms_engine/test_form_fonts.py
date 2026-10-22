"""帳票の字体（ADR-0048）: 中国語の字体の run の言語の書き換えと、字体の置き換えの規則の行き先。"""

import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from django.conf import settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from services import form_docx

ROOT = Path(settings.FORM_TEMPLATES_ROOT)
RULES = Path(settings.BASE_DIR) / "config" / "fonts" / "99-form-fonts.conf"


def _docx(path: Path, runs: list[tuple[str, str, str | None]]) -> Path:
    """（字、東アジアの字体、東アジアの言語）の run を 1 段落に並べた docx。文書の既定は ja-JP。"""
    document = Document()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = OxmlElement("w:lang")
    lang.set(qn("w:eastAsia"), "ja-JP")
    rpr.append(lang)
    paragraph = document.add_paragraph()
    for text, font, language in runs:
        run = paragraph.add_run(text)
        props = run._element.get_or_add_rPr()
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:eastAsia"), font)
        props.append(fonts)
        if language:
            element = OxmlElement("w:lang")
            element.set(qn("w:eastAsia"), language)
            props.append(element)
    document.save(str(path))
    return path


def _languages(path: Path) -> list[str | None]:
    document = Document(str(path))
    result = []
    for run in document.paragraphs[0].runs:
        lang = run._element.rPr.find(qn("w:lang")) if run._element.rPr is not None else None
        result.append(lang.get(qn("w:eastAsia")) if lang is not None else None)
    return result


def test_chinese_font_runs_get_the_chinese_language(tmp_path: Path) -> None:
    source = _docx(
        tmp_path / "source.docx",
        [
            ("有关事项", "SimSun", None),  # 文書の既定（ja-JP）を継ぐ → zh-CN
            ("生活入职教育确认书", "SimHei", "ja-JP"),  # → zh-CN
            ("说明", "SimSun", "zh-TW"),  # → zh-CN（Word は SimSun の字形で描く）
            ("说明", "SimSun", "zh-CN"),  # そのまま
            ("私の日本での生活", "SimSun", None),  # かなを含む → 変えない
            ("男・女", "SimHei", None),  # 中黒はかなとして数えない → zh-CN
            ("說明", "Microsoft JhengHei", "zh-TW"),  # 繁体字の字体 → 変えない
            ("生活", "ＭＳ 明朝", None),  # 日本語の字体 → 変えない
        ],
    )
    copy = tmp_path / "for-pdf.docx"
    fix = form_docx.with_chinese_language(source, copy)
    assert (fix.changed, fix.skipped_kana) == (4, 1)
    assert _languages(copy) == ["zh-CN", "zh-CN", "zh-CN", "zh-CN", None, "zh-CN", "zh-TW", None]
    assert _languages(source)[:3] == [None, "ja-JP", "zh-TW"]  # 保存する文書は変えない


def test_japanese_forms_are_unchanged(tmp_path: Path) -> None:
    """日本語の様式には中国語の字体の run が無いので、写しの中身も変わらない。"""
    source = ROOT / "1-6" / "2026-10-18" / "template.docx"
    copy = tmp_path / "for-pdf.docx"
    fix = form_docx.with_chinese_language(source, copy)
    assert (fix.changed, fix.skipped_kana) == (0, 0)
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(copy) as after:
        assert before.read("word/document.xml") == after.read("word/document.xml")


@pytest.mark.parametrize(
    ("form", "changed"),
    [("1-5", True), ("1-6", True), ("5-8", True), ("5-9", True)],
)
def test_bilingual_forms_have_chinese_runs_to_fix(tmp_path: Path, form: str, changed: bool) -> None:
    """日中併記の様式には、言語が ja-JP・zh-TW の中国語の字体の run がある（ADR-0048 の表）。"""
    source = ROOT / form / "2026-10-18.ja-zh" / "template.docx"
    fix = form_docx.with_chinese_language(source, tmp_path / "for-pdf.docx")
    assert (fix.changed > 0) is changed


def test_simplified_chinese_form_needs_no_fix(tmp_path: Path) -> None:
    """第1-6号（中国語の版）は、全部の run が Noto Serif CJK SC・zh-CN（直す所が無い）。"""
    source = ROOT / "1-6" / "2026-10-18.zh-hans" / "template.docx"
    fix = form_docx.with_chinese_language(source, tmp_path / "for-pdf.docx")
    assert fix.changed == 0


@pytest.mark.skipif(
    shutil.which("fc-match") is None, reason="fontconfig が無い（web のコンテナで流す）"
)
@pytest.mark.parametrize(
    ("requested", "family"),
    [
        ("SimHei", "Noto Sans CJK SC"),
        ("DengXian", "Noto Sans CJK SC"),
        ("Microsoft YaHei", "Noto Sans CJK SC"),
        ("Microsoft JhengHei", "Noto Sans CJK TC"),
        ("Times New Roman", "Liberation Serif"),
        ("Arial", "Liberation Sans"),
        ("Century", "Liberation Serif"),
    ],
)
def test_rules_send_fonts_to_the_installed_families(requested: str, family: str) -> None:
    """規則の行き先がイメージに入っている（fonts-liberation2 などを入れ忘れると別の字体になる）。"""
    assert f"<family>{requested}</family>" in RULES.read_text(encoding="utf-8")
    result = subprocess.run(
        ["fc-match", "-f", "%{family}", requested], capture_output=True, text=True, check=True
    )
    assert family in result.stdout.split(",")
