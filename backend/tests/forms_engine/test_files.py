"""Excel のセルの値だけを書く方式と、Word の差込タグ（制約3。T29、ADR-0018）。

Excel は openpyxl で保存すると図形などが落ちるため、シートの XML のセルの要素だけを書き換える。
書いた後もブックの中のファイルの一覧・図形・印刷の設定・結合セル・入力規則・保護が original と
同じで、ほかのファイルはバイト単位で同じことを確かめる。
"""

import re
import zipfile
from pathlib import Path

import docx
import openpyxl  # type: ignore[import-untyped]
import pytest
from django.conf import settings

from services import form_docx, form_xlsx

ROOT = Path(settings.FORM_TEMPLATES_ROOT)
ORIGINAL_1_17 = ROOT / "1-17" / "2026-10-17" / "original.xlsx"
SHEET = "参考様式第１－１７号"


def structure(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        workbook = archive.read("xl/workbook.xml").decode("utf-8")
        drawing = archive.read("xl/drawings/drawing1.xml")
        return {
            "files": archive.namelist(),
            "drawing": drawing,
            "drawing_ref": re.findall(r"<drawing [^>]*/>", sheet),
            "page_setup": re.findall(r"<pageSetup[^>]*/>", sheet),
            "margins": re.findall(r"<pageMargins[^>]*/>", sheet),
            "print_options": re.findall(r"<printOptions[^>]*/>", sheet),
            "header_footer": re.findall(r"<headerFooter.*?</headerFooter>", sheet, re.S),
            "breaks": re.findall(r"<rowBreaks.*?</rowBreaks>", sheet, re.S),
            "merges": re.findall(r"<mergeCell [^>]*/>", sheet),
            "validations": re.findall(r"<dataValidations.*?</dataValidations>", sheet, re.S),
            "ext": re.findall(r"<extLst>.*?</extLst>", sheet, re.S),
            "protection": re.findall(r"<sheetProtection[^>]*/>", sheet),
            "defined_names": re.findall(r"<definedName .*?</definedName>", workbook),
            "calc": re.findall(r"<calcPr[^>]*/>", workbook),
        }


def test_original_1_17_inventory_is_what_adr_0018_records() -> None:
    """ADR-0018 に書いた中身（変わったら ADR と方式を見直す）。"""
    found = structure(ORIGINAL_1_17)
    files = found["files"]
    assert isinstance(files, list)
    assert "xl/drawings/drawing1.xml" in files and "xl/calcChain.xml" in files
    assert not any("vmlDrawing" in name or "ctrlProps" in name for name in files)
    drawing = found["drawing"]
    assert isinstance(drawing, bytes)
    assert len(re.findall(rb"<xdr:sp[ >]", drawing)) == 5
    assert found["validations"] == [] and found["ext"] == [] and found["protection"] == []
    assert len(found["merges"]) == 843  # type: ignore[arg-type]
    names = found["defined_names"]
    assert isinstance(names, list) and len(names) == 1
    assert "_xlnm.Print_Area" in names[0] and "$A$1:$BP$777" in names[0]


def test_writing_cells_keeps_everything_else(tmp_path: Path) -> None:
    out = tmp_path / "filled.xlsx"
    form_xlsx.write_cells(
        ORIGINAL_1_17,
        out,
        {
            f"{SHEET}!O18": "WANG FANG <テスト> & 名",
            f"{SHEET}!Q26": 1995,
            f"{SHEET}!BC26": "中国",
            f"{SHEET}!CZ5": "outside",  # 行はあるがセルの要素が無い → 列の順に挿入
        },
    )
    assert structure(out) == structure(ORIGINAL_1_17)
    with zipfile.ZipFile(ORIGINAL_1_17) as before, zipfile.ZipFile(out) as after:
        for name in before.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert before.read(name) == after.read(name), name
    sheet = openpyxl.load_workbook(out)[SHEET]  # 書いたセルを読み戻す
    assert sheet["O18"].value == "WANG FANG <テスト> & 名"
    assert sheet["Q26"].value == 1995
    assert sheet["BC26"].value == "中国"
    assert sheet["CZ5"].value == "outside"


@pytest.mark.parametrize(
    "ref",
    [
        "P18",  # 結合セル O18:AJ25 の左上以外
        "DA670",  # 計算式のセル
    ],
)
def test_cells_that_cannot_be_written(tmp_path: Path, ref: str) -> None:
    with pytest.raises(form_xlsx.XlsxWriteError):
        form_xlsx.write_cells(ORIGINAL_1_17, tmp_path / "x.xlsx", {f"{SHEET}!{ref}": "x"})
    with pytest.raises(form_xlsx.XlsxWriteError):
        form_xlsx.check_targets(ORIGINAL_1_17, [f"{SHEET}!{ref}"])


def test_mapping_targets_of_1_17_are_writable() -> None:
    import json

    definition = json.loads(
        (Path(settings.FORM_MAPPINGS_ROOT) / "1-17" / "2026-10-17.json").read_text("utf-8")
    )
    form_xlsx.check_targets(ORIGINAL_1_17, [item["target"] for item in definition["fields"]])


def test_word_template_only_adds_tags_to_the_original(tmp_path: Path) -> None:
    folder = ROOT / "5-9" / "2026-10-17"
    original = docx.Document(str(folder / "original.docx"))
    template = docx.Document(str(folder / "template.docx"))
    original_text = "\n".join(p.text for p in original.paragraphs)
    template_text = "\n".join(p.text for p in template.paragraphs)
    assert form_docx.strip_tags(template_text) == original_text
    assert "{{ organization_name }}" in template_text

    out = tmp_path / "filled.docx"
    form_docx.render(
        folder / "template.docx", out, {"organization_name": "A&B 登録支援機関", "s1_year": 2026}
    )
    filled = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "A&B 登録支援機関" in filled and "2026年" in filled
    assert "{{" not in filled
