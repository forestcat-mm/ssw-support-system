"""支援計画書（参考様式第1-17号）の生成（T31、ADR-0020）。

- 段階3の完了条件 A「様式との一致」①②: ダミーの計画から生成し、頁数と固定の文言が空の様式と同じ。
  書いたセルを読み戻すと、マッピングのすべての値が所定のセルに入っている。
- 制約3: 図形・印刷の設定・結合セルなどが original と同じ。
"""

import datetime
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import pytest
from django.conf import settings
from django.core.exceptions import ValidationError
from django.test import Client
from pytest_django.fixtures import Settings as SettingsWrapper

from apps.forms_engine.models import FormTemplate, GeneratedDocument, GenerationStatus
from apps.support.models import DeliveredBy, DeliveryKind, DeliveryMethod, SupportPlanVersion
from apps.workers.models import ForeignWorker
from services import (
    deliveries,
    form_data_1_17,
    form_generation,
    form_samples,
    form_templates,
    form_xlsx,
    support_plans,
)
from tests.core.factories import make_user, verified_client
from tests.forms_engine import pdfcheck
from tests.forms_engine.conftest import BASE
from tests.forms_engine.test_files import structure

pytestmark = pytest.mark.django_db

ROOT = Path(settings.FORM_TEMPLATES_ROOT)
ORIGINAL = ROOT / "1-17" / "2026-10-17" / "original.xlsx"
MAPPING = Path(settings.FORM_MAPPINGS_ROOT) / "1-17" / "2026-10-17.json"
SHEET = "参考様式第１－１７号"


def generate(worker: ForeignWorker, **kwargs: Any) -> GeneratedDocument:
    target, data = form_samples.sample("1-17", worker, BASE)
    document = form_generation.request(
        "1-17", target=target, data=data, basis_on=BASE, by=None, accept_warnings=True, **kwargs
    )
    document.refresh_from_db()
    return document


# --- マッピング ------------------------------------------------------------------------


def test_mapping_is_the_output_of_the_tool() -> None:
    """マッピングの JSON は tools/forms/make_1_17_mapping.py の出力そのもの（手で直さない）。"""
    tool = Path(settings.REPO_ROOT) / "tools" / "forms" / "make_1_17_mapping.py"
    output = subprocess.run([sys.executable, str(tool)], check=True, capture_output=True).stdout
    assert json.loads(output) == json.loads(MAPPING.read_text(encoding="utf-8"))


def test_printed_texts_in_templates_are_the_original() -> None:
    """印刷済みの文に差し込む欄は、値が空のとき original と同じ文になる（F711・A764 など）。"""
    sheet = openpyxl.load_workbook(ORIGINAL)[SHEET]
    definition = json.loads(MAPPING.read_text(encoding="utf-8"))
    items = {item["target"].split("!")[1]: item for item in definition["fields"]}
    for ref in ("F711", "A764", "BC664", "BC667"):  # Ⅴ・面談の希望・9 の「□ 対面」「□ オンライン」
        item = items[ref]
        empty = {name: spec.get("empty", "") for name, spec in item["keys"].items()}
        assert item["template"].format(**empty) == sheet[ref].value, ref


def test_every_target_is_writable_and_not_referenced_by_formulas() -> None:
    definition = json.loads(MAPPING.read_text(encoding="utf-8"))
    targets = [item["target"] for item in definition["fields"]]
    form_xlsx.check_targets(ORIGINAL, targets)
    assert form_xlsx.targets_referenced_by_formulas(ORIGINAL, targets) == []


# --- 生成（A の①②・制約3）------------------------------------------------------------


def test_generated_from_a_dummy_plan_matches_the_blank_form(seeded: list[ForeignWorker]) -> None:
    worker = seeded[14]  # 15 番（食品。キャリアアッププランの警告のみ）
    document = generate(worker)
    assert document.status == GenerationStatus.COMPLETED, document.failure_reason
    assert isinstance(document.target, SupportPlanVersion)
    snapshot = json.loads(document.snapshot)
    data = snapshot["data"]
    template = FormTemplate.objects.get(form_id="1-17")
    values = form_generation.field_values(template.mapping.definition, data)
    assert len(values) > 100

    # ① 頁数と固定の文言（書いた値と、書き込み先の印刷済みの文字を除いた文字の数が同じ）
    pdf = Path(document.pdf_file.path)
    blank = pdfcheck.blank_pdf(ORIGINAL)
    assert document.page_count == len(pdfcheck.pdf_text(blank))
    sheet = openpyxl.load_workbook(ORIGINAL)[SHEET]
    printed = [sheet[target.split("!")[1]].value for target in values]
    assert pdfcheck.fixed_characters_match(pdf, blank, values.values(), printed)

    # ② 書いたセルを読み戻す
    filled = openpyxl.load_workbook(document.filled_file.path)[SHEET]
    for target, value in values.items():
        assert str(filled[target.split("!")[1]].value) == str(value), target

    # 制約3: 図形・印刷の設定・結合セル・入力規則・保護が original と同じ
    assert structure(Path(document.filled_file.path)) == structure(ORIGINAL)


def test_long_values_are_written_smaller_and_the_size_is_kept(
    seeded: list[ForeignWorker],
) -> None:
    worker = seeded[14]
    document = generate(worker)
    fitted = json.loads(document.snapshot)["fitted"]
    assert fitted and all(6 <= entry["size"] < 13 for entry in fitted.values())
    target = next(iter(fitted))
    with zipfile.ZipFile(document.filled_file.path) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    ref = target.split("!")[1]
    cell = sheet[sheet.index(f'<c r="{ref}"') :].split("</c>", 1)[0]
    assert f'<sz val="{fitted[target]["size"]:g}"/>' in cell and "ＭＳ Ｐ明朝" in cell


def test_text_that_does_not_fit_even_at_the_minimum_is_a_warning(
    seeded: list[ForeignWorker],
) -> None:
    worker = seeded[14]
    version = form_samples.latest_plan_version(worker)
    assert version is not None
    content = support_plans.content_of(version)
    content.activities[
        "1.main"
    ].planned_note = "在留資格認定証明書交付申請前（在留資格変更許可申請前）"
    longer = support_plans.add_version(
        version.plan, content, applies_from=version.applies_from, by=None, change_reason="長い文"
    )
    data = form_data_1_17.build(longer, BASE) | form_data_1_17.declaration_values(True)
    result = form_generation.check("1-17", target=longer, data=data, basis_on=BASE)
    assert [f.message for f in result.warnings if f.code == "text_overflow"] == [
        "「Ⅳ 1.main 実施予定 有」は 6 pt でも欄に収まりません。"
    ]


# --- 生成前の検査（T31 で足した規則）------------------------------------------------------


def test_plan_checks(seeded: list[ForeignWorker]) -> None:
    worker = seeded[14]
    version = form_samples.latest_plan_version(worker)
    assert version is not None
    content = support_plans.content_of(version)
    content.activities["2.a"].outsourced = True
    content.activities["2.a"].outsourcee_name = "（ダミー）送迎会社"
    content.activities["2.a"].outsourcee_address = "東京都サンプル区1-1"
    content.activities["4.main"].methods = [*content.activities["4.main"].methods, "other"]
    content.activities["4.main"].method_notes = {"other": "動画"}
    changed = support_plans.add_version(
        version.plan, content, applies_from=version.applies_from, by=None, change_reason="検査"
    )
    data = form_data_1_17.build(changed, BASE)  # 申告なし
    result = form_generation.check("1-17", target=changed, data=data, basis_on=BASE)
    codes = {finding.code: finding for finding in result.findings}
    assert codes["plan_outsourced"].severity == "block"
    assert codes["orientation_other"].severity == "warn"
    assert codes["propriety_not_declared"].severity == "warn"


def test_contract_past_its_end_date_is_a_warning(seeded: list[ForeignWorker]) -> None:
    """期間の終了日を過ぎて終了を登録していない契約（10 番。D+20 日に満了）。"""
    worker = seeded[9]
    later = BASE + datetime.timedelta(days=21)
    target, data = form_samples.sample("1-17", worker, later)
    result = form_generation.check("1-17", target=target, data=data, basis_on=later)
    [finding] = [f for f in result.warnings if f.code == "contract_past_end"]
    ends_on = (BASE + datetime.timedelta(days=20)).isoformat()
    assert finding.message == (
        f"雇用契約の期間の終了日（{ends_on}）を過ぎています。更新の新しい版か終了を登録してください。"
    )
    assert finding.destination is not None
    assert finding.destination.url() == f"/workers/{worker.pk}/contracts/"


def test_builder_parts() -> None:
    assert form_data_1_17.postal_parts("000-0001") == ("000", "0001")
    assert form_data_1_17.postal_parts("〒1000001") == ("100", "0001")
    assert form_data_1_17.registration_parts("99登-999999") == ("99", "999999")
    assert form_data_1_17.registration_parts("19登－000123") == ("19", "000123")
    assert form_data_1_17.hour_text("09:00") == "9時"
    assert form_data_1_17.hour_text("17:30") == "17時30分"


# --- 開発の DB での見本の版の置き換え（T31-D1）------------------------------------------


@pytest.mark.parametrize(
    ("debug", "by_user", "replaced"),
    [(False, False, False), (True, False, True), (True, True, False)],
)
def test_sample_versions_are_replaced_only_in_debug(
    repo_copy: tuple[Path, Path],
    settings: SettingsWrapper,
    seeded: list[ForeignWorker],
    debug: bool,
    by_user: bool,
    replaced: bool,
) -> None:
    worker = seeded[14]
    target, data = form_samples.sample("1-17", worker, BASE)
    user = make_user("kanri", roles=("admin",)) if by_user else None
    form_generation.request(
        "1-17", target=target, data=data, basis_on=BASE, by=user, accept_warnings=True
    )
    _templates, mappings = repo_copy
    path = mappings / "1-17" / "2026-10-17.json"
    definition = json.loads(path.read_text(encoding="utf-8"))
    definition["note"] = "changed"
    path.write_text(json.dumps(definition, ensure_ascii=False), encoding="utf-8")
    settings.DEBUG = debug
    if replaced:
        assert form_templates.load_all().updated == ["1-17 2026-10-17 ja"]
    else:
        with pytest.raises(form_templates.FormTemplateError, match="生成した履歴があります"):
            form_templates.load_all()


# --- 交付記録（T31-D16）------------------------------------------------------------------


def delivery_input(document: GeneratedDocument | None, on: datetime.date) -> Any:
    return deliveries.DeliveryInput(
        kind=DeliveryKind.SUPPORT_PLAN_COPY,
        delivered_on=on,
        method=DeliveryMethod.HAND,
        delivered_by=DeliveredBy.ORGANIZATION,
        explained=True,
        generated_document=document,
    )


def test_delivery_from_a_generated_document_takes_its_plan_version(
    seeded: list[ForeignWorker],
) -> None:
    worker = seeded[14]
    document = generate(worker)
    version = document.target
    assert isinstance(version, SupportPlanVersion)
    # 新しい版（適用開始日の前）を作り、その写しを今日渡す → 版は生成物の版で、違いを警告する
    newer = support_plans.add_version(
        version.plan,
        support_plans.content_of(version),
        applies_from=BASE + datetime.timedelta(days=30),
        by=None,
        change_reason="次の版",
    )
    target, data = form_samples.sample("1-17", worker, BASE)
    assert target == newer
    newer_document = generate(worker)
    delivery = deliveries.create_delivery(
        worker, delivery_input(newer_document, BASE), by=None, today=BASE
    )
    assert delivery.plan_version == newer and delivery.generated_document == newer_document
    [warning] = [w for w in deliveries.warnings(delivery) if w.code == "version_differs"]
    assert str(newer.version) in warning.message
    # 手で作る交付記録は、今のとおり交付日の版
    manual = deliveries.create_delivery(worker, delivery_input(None, BASE), by=None, today=BASE)
    assert manual.plan_version == version and manual.generated_document is None


def test_delivery_only_from_this_workers_1_17(seeded: list[ForeignWorker]) -> None:
    document = generate(seeded[14])
    with pytest.raises(ValidationError) as error:
        deliveries.create_delivery(
            seeded[11], delivery_input(document, BASE), by=None, today=BASE
        )  # 12 番（別の人）
    assert error.value.error_dict["__all__"][0].code == "generated_document"


# --- 画面 ------------------------------------------------------------------------------


def test_generation_screen_document_screen_and_delivery(seeded: list[ForeignWorker]) -> None:
    worker = seeded[14]
    version = form_samples.latest_plan_version(worker)
    assert version is not None
    client = verified_client(make_user("tanto9", roles=("support_staff",)))
    url = (
        f"/workers/{worker.pk}/support/plans/{version.plan_id}/versions/{version.version}/generate/"
    )
    page = client.get(url)
    assert page.status_code == 200
    assert "キャリアアッププランの交付の記録がありません" in page.content.decode()
    form = {"basis_on": "2026-10-16", "propriety": "on", "action": "generate"}
    once = page.context["once_token"]
    response = client.post(url, form | {page.context["once_field"]: once})
    assert response.status_code == 200  # 警告を確かめていない
    assert "警告を確かめてから生成してください" in response.content.decode()
    once = response.context["once_token"]  # 表示し直すと新しい札になる（二重送信の防止）
    response = client.post(url, form | {"accept_warnings": "on", page.context["once_field"]: once})
    document = GeneratedDocument.objects.get(worker=worker)
    assert (
        response.status_code == 302 and response["Location"] == f"/forms/documents/{document.pk}/"
    )
    detail = client.get(response["Location"])
    assert detail.status_code == 200
    content = detail.content.decode()
    assert f"/forms/documents/{document.pk}/pdf/" in content
    delivery_url = (
        f"/workers/{worker.pk}/support/deliveries/new/support_plan_copy/?generated={document.pk}"
    )
    assert delivery_url.replace("&", "&amp;") in content or delivery_url in content
    form_page = client.get(delivery_url)
    assert form_page.status_code == 200 and 'id="delivery-generated"' in form_page.content.decode()
    assert client.get(f"/forms/documents/{document.pk}/pdf/").status_code == 200

    no_forms = make_user("noforms2", roles=("support_staff",))
    no_forms.user_roles.get().role.permissions.filter(action="forms.generate").delete()
    assert verified_client(no_forms).get(url).status_code == 403
    assert verified_client(no_forms).get(f"/forms/documents/{document.pk}/").status_code == 403
    assert Client().get(url).status_code in (302, 403)


def test_generation_screen_when_the_form_is_not_loaded(settings: SettingsWrapper) -> None:
    """様式を読み込む前（make forms の前）に開いても 500 にせず、読み込みを頼む文を出す。"""
    from tests.contracts.factories import make_contract
    from tests.workers.factories import make_worker

    worker = make_worker()
    contract = make_contract(worker)
    plan = support_plans.create_plan(worker, contract.host_organization, applies_from=BASE, by=None)
    version = plan.versions.get()
    client = verified_client(make_user("tanto8", roles=("support_staff",)))
    url = f"/workers/{worker.pk}/support/plans/{plan.pk}/versions/{version.version}/generate/"
    response = client.get(url)
    assert response.status_code == 200
    assert "様式が読み込まれていません" in response.content.decode()
