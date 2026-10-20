"""帳簿の出力（T57。ADR-0036・ADR-0044）。

外国人ごとの帳簿（③④）を画面と PDF で出せ、機関全体の xlsx に①②と月ごとの集計が出る。
出力は監査ログに EXPORT。
"""

import datetime
import io
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import pytest
from django.utils import timezone

from apps.core.models import AuditAction, AuditLog
from apps.filings.models import FilingKind, Trigger
from apps.hosts.models import SupportContract
from services import book_exports, filings
from tests.core.factories import make_user, verified_client
from tests.forms_engine.test_form_application import make_case
from tests.support.test_record_screens import make_record
from tests.workers.factories import add_residence

pytestmark = pytest.mark.django_db


@pytest.fixture
def case() -> Any:
    contract = make_case()
    SupportContract.objects.filter(host_organization=contract.host_organization).update(
        ends_on=datetime.date(2027, 3, 31)
    )
    worker = contract.worker
    add_residence(worker)
    from tests.core.factories import make_staff

    worker.assigned_staff = make_staff(name="佐藤 花子", position="支援担当者")
    worker.save()
    make_record(
        worker,
        occurred_on=datetime.date(2026, 6, 1),
        details={"event": "夏祭り", "place": "公民館"},
    )
    return contract


def test_worker_book_lists_the_records(case: Any) -> None:
    worker = case.worker
    book = book_exports.worker_book(worker)
    assert book.card_number == "AB12345678CD"
    assert book.hosts == [case.host_organization.name]
    [period] = book.periods
    assert period.starts_on == datetime.date(2026, 4, 1)
    [record] = book.records
    assert (record.place, record.staff, record.position) == ("公民館", "佐藤 花子", "支援担当者")
    sheets = openpyxl.load_workbook(io.BytesIO(book_exports.worker_book_xlsx(book)))
    assert sheets.sheetnames == ["③支援対象者", "④支援の実施", "④関係する書類"]
    assert sheets["④支援の実施"]["D2"].value == "公民館"


def test_monthly_counts(case: Any) -> None:
    filings.create(
        FilingKind.MISSING,
        target=case.worker,
        basis_on=datetime.date(2026, 7, 15),
        trigger=Trigger.MANUAL,
    )
    counts = {
        f"{count.month:%Y-%m}": count
        for count in book_exports.monthly_counts(
            datetime.date(2026, 3, 1), datetime.date(2026, 7, 31)
        )
    }
    assert counts["2026-03"].supported == 0  # 雇用契約は 4 月から
    assert counts["2026-04"].supported == 1
    assert (counts["2026-07"].missing, counts["2026-06"].missing) == (1, 0)
    # 職員数は今の有効な担当者を登録の日から数える（推定）。今月には担当者がいる
    today = timezone.localdate()
    [now] = book_exports.monthly_counts(today, today)
    assert now.full_time + now.part_time >= 1


def test_screens_and_exports(case: Any) -> None:
    worker = case.worker
    client = verified_client(make_user("kanri-books", roles=("admin",)))
    page = client.get(f"/workers/{worker.pk}/books/").content.decode()
    assert 'id="book-records"' in page and "公民館" in page
    pdf = client.get(f"/workers/{worker.pk}/books/pdf/")
    assert pdf.status_code == 200
    assert b"".join(pdf.streaming_content)[:4] == b"%PDF"  # type: ignore[attr-defined]
    assert AuditLog.objects.filter(
        action=AuditAction.EXPORT, object_id=str(worker.pk), context__books="worker"
    ).exists()
    form = client.get("/settings/books/")
    assert form.status_code == 200
    response = client.post("/settings/books/", {"starts_on": "2026-04-01", "ends_on": "2026-10-31"})
    assert response.status_code == 200
    content = b"".join(response.streaming_content)  # type: ignore[attr-defined]
    book = openpyxl.load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["①支援実施体制", "①担当者", "②支援の委託契約", "月ごとの集計"]
    assert book["月ごとの集計"]["A2"].value == "2026-04"
    assert AuditLog.objects.filter(
        action=AuditAction.EXPORT, context__books="organization"
    ).exists()
    nobody = verified_client(make_user("nobody-books"))
    assert nobody.get("/settings/books/").status_code == 403
